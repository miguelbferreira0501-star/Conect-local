"""Matriz de decisão: notas 0-10 por critério, pesos da missão, nota final 0-100.
Gera matriz.xlsx e matriz.md. Rodar: python3 matriz.py
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

CRIT = [
    ("Gravidade/frequência", 0.15),
    ("Disposição a pagar", 0.15),
    ("Evidência de demanda", 0.12),
    ("1ºs clientes", 0.10),
    ("MVP viável", 0.10),
    ("Receita e margem", 0.10),
    ("Retenção", 0.08),
    ("Espaço/diferenciação", 0.08),
    ("Vantagem IA/tec", 0.05),
    ("Expansão", 0.04),
    ("Risco (alto=baixo risco)", 0.03),
]

# (id, nome, público, modelo de receita, notas[11], confiança A/M/B)
IDEIAS = [
    (1, "Pasta em Dia", "Escritórios contábeis de 1-10 pessoas (30-300 CNPJs)", "Assinatura por escritório, faixa por nº de clientes", [8, 8, 8, 8, 7, 7, 9, 6, 7, 8, 7], "M"),
    (2, "Raio-X da Reforma", "Escritórios contábeis com clientes no Simples", "Assinatura sazonal / por diagnóstico", [7, 6, 7, 7, 6, 6, 4, 5, 6, 5, 6], "M"),
    (3, "Orçamento na Hora (oficina)", "Oficinas mecânicas independentes 1-5 elevadores", "Assinatura mensal", [8, 5, 6, 7, 6, 5, 7, 5, 6, 7, 7], "B"),
    (4, "Pedido Único (representante)", "Representantes com 3-10 representadas", "Assinatura por representante", [7, 7, 5, 6, 4, 6, 8, 7, 7, 6, 5], "B"),
    (5, "Comissão Certa", "Representantes comerciais autônomos", "Assinatura ou % do recuperado", [5, 6, 3, 6, 6, 5, 7, 7, 5, 5, 7], "B"),
    (6, "Base Solar (O&M)", "Integradores solares pequenos (base de 50-1.000 usinas)", "Assinatura por integrador + por usina", [7, 7, 7, 7, 5, 7, 8, 6, 4, 6, 6], "M"),
    (7, "Homologa (solar)", "Integradores solares", "Assinatura", [7, 6, 7, 7, 6, 5, 5, 7, 4, 4, 7], "M"),
    (8, "Extra Aprovado (obra)", "Empreiteiros e empresas de reforma", "Assinatura barata", [7, 4, 6, 5, 8, 4, 5, 7, 4, 6, 7], "B"),
    (9, "Aluguel Direto", "Locadores sem imobiliária e imobiliárias pequenas", "Assinatura por imóvel", [7, 6, 7, 5, 6, 6, 9, 5, 4, 7, 5], "M"),
    (10, "Vistoria em Vídeo", "Imobiliárias pequenas e locadores", "Por vistoria", [6, 5, 5, 6, 6, 4, 3, 4, 7, 5, 7], "B"),
    (11, "Carteira do Síndico", "Síndicos profissionais (5-20 condomínios)", "Assinatura por síndico", [6, 6, 5, 7, 6, 6, 8, 5, 5, 5, 8], "B"),
    (12, "Canhoto Digital", "Transportadoras pequenas", "Por entrega / assinatura", [6, 6, 5, 5, 5, 5, 8, 4, 6, 5, 7], "B"),
    (13, "Repasse Conferido (iFood)", "Restaurantes com >R$ 30 mil/mês no iFood", "Assinatura ou % do recuperado", [8, 8, 7, 7, 5, 7, 8, 6, 6, 6, 5], "M"),
    (14, "Margem Real (ML/Shopee)", "Vendedores de marketplace 100-3.000 pedidos/mês", "Assinatura", [7, 6, 6, 5, 6, 5, 7, 4, 5, 7, 6], "M"),
    (15, "Sentinela de Integração", "Lojistas multicanal e agências", "Assinatura", [7, 6, 6, 5, 4, 6, 7, 6, 4, 6, 5], "M"),
    (16, "Zap Justo", "Empresas com 2-20 atendentes no WhatsApp", "Assinatura + custo Meta repassado", [7, 7, 8, 4, 5, 6, 8, 3, 5, 7, 4], "M"),
    (17, "Prontuário Livre", "Psicólogos, nutris e fisios solo", "Assinatura", [6, 6, 6, 5, 5, 5, 9, 3, 5, 6, 4], "M"),
    (18, "Atendente de nicho no WhatsApp", "Ar-condicionado, dedetização, autoescola", "Assinatura", [7, 6, 6, 6, 6, 6, 7, 3, 8, 7, 5], "M"),
    (19, "Orçamento pela Foto", "Vidraçarias, serralherias, marcenarias", "Assinatura", [6, 5, 4, 6, 5, 5, 6, 7, 8, 6, 6], "B"),
    (20, "Prontuário Vet por Voz", "Clínicas veterinárias pequenas", "Freemium + assinatura", [7, 6, 5, 6, 5, 6, 8, 6, 8, 6, 6], "B"),
    (21, "Pede Avaliação", "Prestadores locais (qualquer)", "Assinatura barata", [5, 5, 6, 7, 8, 4, 6, 3, 4, 6, 8], "M"),
    (22, "Folga Certa (fim da 6x1)", "Pequeno varejo 3-30 funcionários", "Diagnóstico avulso", [7, 3, 7, 6, 7, 4, 4, 2, 3, 4, 6], "A"),
    (23, "NR-1 para micro", "Empresas de 5-50 funcionários", "Assinatura / laudo", [6, 5, 7, 5, 6, 5, 6, 2, 4, 5, 5], "A"),
    (24, "Rimou (batalhas de rima)", "Organizadores de batalha, MCs, contratantes", "Comissão de shows + patrocínio", [5, 3, 5, 6, 7, 4, 6, 8, 2, 7, 6], "M"),
    (25, "Depois (pós-óbito via funerária)", "Funerárias e planos funerários", "Por família ou por plano", [7, 4, 5, 5, 7, 4, 3, 7, 4, 5, 6], "M"),
    (26, "Integra (imigrantes)", "Construtoras e frigoríficos com imigrantes", "Por turma / mensal", [7, 4, 4, 4, 5, 4, 6, 7, 5, 4, 4], "M"),
    (27, "Mensalidade Automática (cursos)", "Escolas de idiomas, música, esportes", "Assinatura + taxa por cobrança", [7, 6, 6, 6, 6, 5, 8, 4, 3, 6, 6], "M"),
    (28, "Pataí (referência: app atual)", "Pet shops de banho e tosa", "Assinatura R$ 49,99", [6, 4, 6, 5, 10, 4, 7, 3, 2, 6, 8], "M"),
]

RUBRICA = {
    "Gravidade/frequência": "9: diário/semanal, perde dinheiro mensurável, gambiarra documentada · 7: mensal, custa horas · 5: real mas ocasional · 3: conveniência",
    "Disposição a pagar": "9: já pagam concorrentes >R$100/mês e o ROI é claro · 7: pagam por algo vizinho · 5: incerto · 3: público sem orçamento",
    "Evidência de demanda": "9: várias fontes primárias + concorrentes com receita · 7: 2-3 fontes + concorrentes · 5: sinais de fornecedor · 3: só hipótese",
    "1ºs clientes": "9: comprador concentrado, alcançável por WhatsApp/visita, ciclo de dias · 5: precisa de vários contatos · 3: venda corporativa lenta ou depende de viralizar",
    "MVP viável": "9: semanas, sem integrações frágeis · 7: 4-8 semanas com 1 integração · 5: 1 integração difícil · 3: várias integrações ou regulação pesada",
    "Receita e margem": "9: >R$300/mês por cliente e margem >80% · 7: R$150-300 · 5: R$50-150 · 3: <R$50 ou margem baixa",
    "Retenção": "9: uso obrigatório todo mês, troca dolorosa · 7: uso mensal · 5: sazonal · 3: uso único",
    "Espaço/diferenciação": "9: ninguém resolve bem, cunha clara · 7: concorrentes com lacuna documentada · 5: concorrentes razoáveis · 3: lotado",
    "Vantagem IA/tec": "9: IA torna viável algo antes caro · 7: IA melhora muito · 5: ajuda · 3: irrelevante",
    "Expansão": "9: segmentos vizinhos e outros países · 7: nacional + vizinhos · 5: nacional · 3: local",
    "Risco (alto=baixo risco)": "9: sem regulação nem dependência · 5: depende de terceiro (Meta, iFood) ou dado sensível · 3: regulação pesada",
}


def nota(scores):
    return round(sum(s * w for s, (_, w) in zip(scores, CRIT)) * 10, 1)


def main():
    rows = sorted(IDEIAS, key=lambda r: -nota(r[4]))
    wb = Workbook()
    ws = wb.active
    ws.title = "Matriz"
    head = ["Rank", "#", "Ideia", "Público", "Receita"] + [f"{c} ({int(w*100)}%)" for c, w in CRIT] + ["Nota (0-100)", "Confiança"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="9D2B5A")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, (id_, nome, pub, rec, sc, conf) in enumerate(rows, start=2):
        ws.append([i - 1, id_, nome, pub, rec] + sc + [None, conf])
        # fórmula na planilha, para conferir e mudar notas
        cols = [get_column_letter(6 + k) for k in range(len(CRIT))]
        ws.cell(row=i, column=6 + len(CRIT)).value = "=(" + "+".join(f"{col}{i}*{w}" for col, (_, w) in zip(cols, CRIT)) + ")*10"
    widths = [6, 4, 30, 38, 30] + [12] * len(CRIT) + [12, 10]
    for k, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(k)].width = w
    ws.freeze_panes = "F2"
    r = wb.create_sheet("Régua")
    r.append(["Critério", "Peso", "O que vale cada nota"])
    for c, w in CRIT:
        r.append([c, w, RUBRICA[c]])
    r.column_dimensions["A"].width = 26
    r.column_dimensions["C"].width = 140
    wb.save("matriz.xlsx")

    lines = ["| Rank | Ideia | Público | Receita | Nota | Confiança |", "|---|---|---|---|---|---|"]
    for i, (id_, nome, pub, rec, sc, conf) in enumerate(rows, start=1):
        lines.append(f"| {i} | {nome} (#{id_}) | {pub} | {rec} | **{nota(sc)}** | {conf} |")
    open("matriz.md", "w").write("\n".join(lines) + "\n")
    for i, r_ in enumerate(rows[:8], 1):
        print(i, r_[1], nota(r_[4]))


if __name__ == "__main__":
    main()
